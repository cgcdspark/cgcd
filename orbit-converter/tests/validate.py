"""Inspects a converted file with tools that share no code with the app. Prints JSON."""
import io
import json
import os
import subprocess
import sys
import tempfile
import zipfile

import docx as pydocx
import openpyxl

from PIL import Image


def image_info(data):
    im = Image.open(io.BytesIO(data))
    fmt = im.format
    im.load()  # full decode: catches truncated or corrupt pixel data
    return {"format": fmt, "width": im.width, "height": im.height, "mode": im.mode}


def pdf_info(path):
    info = subprocess.run(["pdfinfo", path], capture_output=True, text=True)
    if info.returncode != 0:
        return {"format": None, "error": info.stderr.strip()}
    pages = int(next(l.split()[-1] for l in info.stdout.splitlines() if l.startswith("Pages:")))
    text = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
    return {"format": "PDF", "pages": pages, "text": text}


def office_opens(path):
    """True when LibreOffice, an independent office suite, can open the file and print it to PDF."""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, path], capture_output=True, timeout=180)
        pdfs = [f for f in os.listdir(tmp) if f.endswith(".pdf")]
        return bool(pdfs) and os.path.getsize(os.path.join(tmp, pdfs[0])) > 0


def docx_info(path):
    d = pydocx.Document(path)
    paras = [{"text": p.text, "style": p.style.name if p.style is not None else None} for p in d.paragraphs]
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    images = sum(1 for r in d.inline_shapes)
    landscape = d.sections[0].page_width > d.sections[0].page_height
    return {"format": "DOCX", "paragraphs": paras, "tables": tables, "images": images, "landscape": landscape}


def xlsx_info(path):
    wb = openpyxl.load_workbook(path)
    return {"format": "XLSX", "sheets": {ws.title: [[c for c in row] for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}}


def main(path):
    data = open(path, "rb").read()
    if data[:5] == b"%PDF-":
        return pdf_info(path)
    if data[:4] == b"PK\x03\x04":
        z = zipfile.ZipFile(io.BytesIO(data))
        names = z.namelist()
        if "word/document.xml" in names:
            return docx_info(path)
        if "xl/workbook.xml" in names:
            return xlsx_info(path)
        entries = []
        for n in names:
            raw = z.read(n)
            try:
                entries.append({"name": n, **image_info(raw)})
            except Exception:
                entries.append({"name": n, "format": "TEXT", "text": raw.decode("utf-8")})
        return {"format": "ZIP", "entries": entries}
    try:
        return image_info(data)
    except Exception:
        pass
    try:
        return {"format": "TEXT", "text": data.decode("utf-8")}
    except UnicodeDecodeError:
        return {"format": None}


result = main(sys.argv[1])
if "--office" in sys.argv:
    result["office_opens"] = office_opens(sys.argv[1])
print(json.dumps(result, default=str))
